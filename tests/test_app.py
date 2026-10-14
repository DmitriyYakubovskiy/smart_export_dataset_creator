import io
import importlib
import os
import tempfile
from pathlib import Path
import sys

import pytest
from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
WORK = tempfile.TemporaryDirectory()
os.environ['DATA_DIR'] = WORK.name
os.environ['APP_PASSWORD'] = 'test-password-123456'
from app import main

HEADERS={'X-Requested-With':'Anamnesis'}

def file_bytes():
    doc=Document();doc.add_paragraph('Анамнез заболевания: текст пациента <script>alert(1)</script>')
    table=doc.add_table(rows=1,cols=3);table.cell(0,0).merge(table.cell(0,2)).text='Объединённая ячейка'
    buf=io.BytesIO();doc.save(buf);return buf.getvalue()

@pytest.fixture
def client():
    c=TestClient(main.app,headers=HEADERS)
    assert c.post('/api/login',json={'username':'admin','password':'test-password-123456'}).status_code==200
    return c

def upload(c):
    response=c.post('/api/documents',files={'file':('история.docx',file_bytes())})
    assert response.status_code==201
    return response.json()['id']

def test_auth_and_csrf():
    c=TestClient(main.app)
    assert c.get('/api/documents').status_code==401
    assert c.post('/api/login',json={'username':'admin','password':'test-password-123456'}).status_code==403
    assert c.post('/api/login',headers=HEADERS,json={'username':'admin','password':'bad'}).status_code==401

def test_document_preview_and_protection(client):
    identifier=upload(client)
    text=client.get(f'/api/documents/{identifier}/preview').text
    assert text.count('Объединённая ячейка')==1
    assert '&lt;script&gt;' in text and '<script>alert(1)</script>' not in text
    response = client.get(f'/api/documents/{identifier}/download')
    assert response.status_code == 200
    assert response.content == (main.DATA / 'files' / f'{identifier}.docx').read_bytes()
    anonymous=TestClient(main.app)
    assert anonymous.get(f'/api/documents/{identifier}/preview').status_code==401
    assert anonymous.get(f'/api/documents/{identifier}/download').status_code==401

def test_crud_export_persistence(client):
    identifier=upload(client)
    response=client.post(f'/api/documents/{identifier}/histories',json={'title':'Заболевание','text':'Первый анамнез'})
    assert response.status_code==201
    hid=response.json()['id']
    client.post(f'/api/documents/{identifier}/histories',json={'text':'Второй анамнез'})
    assert len(client.get(f'/api/documents/{identifier}/histories').json())==2
    assert client.put(f'/api/histories/{hid}',json={'title':'Новое название','text':'Уточнённый анамнез'}).status_code==200
    with main.db() as c:
        row=c.execute('SELECT * FROM histories WHERE id=?',(hid,)).fetchone()
        assert row['document_id']==identifier and row['text']=='Уточнённый анамнез'
    rows=client.get('/api/export/json').json()
    row=next(r for r in rows if r['id']==hid)
    assert row['document_name']=='история.docx' and row['text']=='Уточнённый анамнез'
    assert client.get('/api/export/csv').content.startswith(b'\xef\xbb\xbf')
    for h in client.get(f'/api/documents/{identifier}/histories').json():
        columns=client.get('/api/labels').json()
        entered={l['id']:(0 if l['data_type']=='int' else False if l['data_type']=='bool' else l['options'][0]) for l in columns if l['active']}
        assert client.put('/api/histories/'+h['id'],json={'title':h['title'],'text':h['text'],'values':entered}).status_code==200
    assert client.patch(f'/api/documents/{identifier}',json={'reviewed':True}).status_code==200
    assert next(d for d in client.get('/api/documents').json() if d['id']==identifier)['reviewed']==1
    assert client.delete(f'/api/histories/{hid}').status_code==200
    assert len(client.get(f'/api/documents/{identifier}/histories').json())==1

def test_validation(client):
    assert client.post('/api/documents',files={'file':('bad.docx',b'not a docx')}).status_code==400
    assert client.post('/api/documents',files={'file':('bad.exe',b'data')}).status_code==400
    identifier=upload(client)
    assert client.post(f'/api/documents/{identifier}/histories',json={'text':'   '}).status_code==400
    assert client.post('/api/documents/missing/histories',json={'text':'hello'}).status_code==404

def test_logout(client):
    assert client.post('/api/logout').status_code==200
    assert client.get('/api/me').status_code==401
